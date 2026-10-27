r"""시연 계정별 테스트용 이력서 파일(pdf·docx·hwpx)을 만듭니다. 모두 합성 값입니다.

대화에 서류를 올려 판독(agent/input/documents.py)을 확인할 때 씁니다. 만든 파일은
data/samples/resumes/에 커밋되어 있습니다. 마지막 경력의 퇴사월은 seed의 이직일(입장일 기준
상대값)에 맞춰 만들므로, 시연하는 달이 바뀌면 그날로 다시 만듭니다.

    uv run --with reportlab --with python-hwpx \
        python -m scripts.make_sample_resumes --today 2026-10-27

--today를 생략하면 오늘(한국 시간)입니다.

PDF에는 한글 글꼴이 필요합니다(기본: Windows 맑은 고딕, 다른 경로는 --font). reportlab·python-hwpx는
이 스크립트에서만 써서 프로젝트 의존성에 넣지 않았습니다.
"""

import argparse
from dataclasses import dataclass, field, replace
from datetime import date
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml.ns import qn

from app.core.clock import today_kst
from app.services.seed import at, read_seed_file

OUT_DIR = Path(__file__).resolve().parents[1] / "data" / "samples" / "resumes"
DEFAULT_FONT = Path("C:/Windows/Fonts/malgun.ttf")
DOCX_FONT = "맑은 고딕"
# 본문에 쓰면 AI가 이력서 칸(기타 희망사항)에 옮겨 적어서 파일 속성(설명·주제)에만 둡니다.
NOTICE = "테스트용 합성 이력서입니다. 실제 인물·회사가 아닙니다."
# 마지막 경력의 퇴사월 자리. 만들 때 seed의 이직일(--today 기준)로 채웁니다.
SEP = "{SEP}"


@dataclass
class Resume:
    persona: str
    title: str
    name: str
    birth: str
    phone: str
    email: str
    education: list[tuple[str, str, str]]  # 기간, 학교, 졸업 여부
    careers: list[tuple[str, str, str, str]]  # 기간, 회사, 직위, 하던 일
    certificates: list[tuple[str, str, str]]  # 이름, 발급 기관, 취득일
    trainings: list[tuple[str, str]] = field(default_factory=list)  # 기간, 과정
    wishes: list[tuple[str, str]] = field(default_factory=list)  # 항목, 값


RESUMES = [
    Resume(
        persona="park",
        title="생산직 프레스 경력 이력서",
        name="박진수",
        birth="1965년",
        phone="010-0000-0001",
        email="park.demo@example.com",
        education=[("1981.03 ~ 1984.02", "OO공업고등학교 기계과", "졸업")],
        careers=[
            (f"2004.03 ~ {SEP}", "OO산업", "반장", "프레스 기계 운전, 금형 교체, 생산 품질 검사"),
            ("1990.03 ~ 2004.02", "OO기계", "사원", "부품 가공, 선반 작업"),
        ],
        certificates=[("지게차운전기능사", "한국산업인력공단", "2011.05.20")],
        trainings=[("2009.03 ~ 2009.06", "용접 기초 과정")],
        wishes=[
            ("희망 직종", "생산·제조 (프레스·기계 조작)"),
            ("희망 지역", "서울 OO구"),
            ("고용 형태", "정해지지 않은 일(정규직)"),
        ],  # fmt: skip
    ),
    Resume(
        persona="choi",
        title="웹 개발 경력 이력서 (2026년 고침)",
        name="최영수",
        birth="1959년",
        phone="010-0000-0002",
        email="choi.demo@example.com",
        education=[("1978.03 ~ 1985.02", "OO대학교 전자계산학과", "졸업")],
        careers=[
            (f"2022.07 ~ {SEP}", "OO에이전시", "책임", "기업 홈페이지 개발·유지보수"),
            ("2005.03 ~ 2019.04", "OO웹솔루션", "과장", "웹사이트 개발·유지보수"),
        ],
        certificates=[
            ("정보처리기사", "한국산업인력공단", "1998.11.15"),
            ("웹디자인기능사", "한국산업인력공단", "2021.06.18"),
        ],
        wishes=[
            ("희망 직종", "웹 개발자"),
            ("희망 지역", "서울 노원구"),
            ("고용 형태", "상관없음"),
        ],  # fmt: skip
    ),
    Resume(
        persona="lee",
        title="앱 개발 경력 이력서",
        name="이정호",
        birth="1968년",
        phone="010-0000-0003",
        email="lee.demo@example.com",
        education=[("1987.03 ~ 1994.02", "OO대학교 컴퓨터공학과", "졸업")],
        careers=[
            (f"2012.03 ~ {SEP}", "OO모바일", "팀장", "모바일 앱 개발·운영"),
            ("2001.04 ~ 2012.02", "OO시스템", "대리", "웹·SI 개발"),
        ],
        certificates=[("정보처리기사", "한국산업인력공단", "2003.05.23")],
        wishes=[
            ("희망 직종", "모바일 앱 개발자"),
            ("희망 지역", "서울 강서구"),
            ("고용 형태", "정해지지 않은 일(정규직)"),
        ],  # fmt: skip
    ),
]


def with_separation(r: Resume, separation: date) -> Resume:
    """마지막 경력의 퇴사월({SEP})을 그 계정의 이직월로 채웁니다(seed와 같은 기준)."""
    month = f"{separation:%Y.%m}"
    careers = [(period.replace(SEP, month), *rest) for period, *rest in r.careers]
    return replace(r, careers=careers)


def sections(r: Resume) -> list[tuple[str, list[str], list[list[str]]]]:
    """(소제목, 머리 칸, 행들). 세 형식이 같은 내용을 쓰도록 한 곳에서 만듭니다."""
    return [
        ("인적 사항", ["항목", "내용"],
         [["성명", r.name], ["출생", r.birth], ["연락처", r.phone], ["이메일", r.email]]),
        ("학력", ["기간", "학교", "졸업 여부"], [list(e) for e in r.education]),
        ("경력", ["기간", "회사", "직위", "하던 일"], [list(c) for c in r.careers]),
        ("자격증", ["자격증", "발급 기관", "취득일"], [list(c) for c in r.certificates]),
        ("직업훈련", ["기간", "과정"], [list(t) for t in r.trainings]),
        ("희망 근무 조건", ["항목", "내용"], [list(w) for w in r.wishes]),
    ]  # fmt: skip


def _use_korean_font(style) -> None:
    """기본 서식은 한글 글꼴을 테마로만 정해 두어(제목은 비어 있음) 보는 프로그램에 따라 한글이
    네모로 나옵니다. 모든 글자 칸에 맑은 고딕을 직접 지정합니다."""
    _set_fonts(style.element.get_or_add_rPr().get_or_add_rFonts())


def _set_fonts(fonts) -> None:
    for theme in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(theme), None)
    for slot in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(slot), DOCX_FONT)


def write_docx(r: Resume, path: Path) -> None:
    doc = Document()
    for name in ("Normal", "Heading 1", "Heading 2", "Table Grid"):
        _use_korean_font(doc.styles[name])
    _set_fonts(doc.styles.element.find(qn("w:docDefaults")).find(".//" + qn("w:rFonts")))
    doc.add_heading(f"이력서 - {r.title}", level=1)
    for heading, head, rows in sections(r):
        if not rows:
            continue
        doc.add_heading(heading, level=2)
        table = doc.add_table(rows=1 + len(rows), cols=len(head))
        table.style = "Table Grid"
        for i, row in enumerate([head, *rows]):
            for j, value in enumerate(row):
                table.cell(i, j).text = value
    doc.core_properties.comments = NOTICE
    doc.save(str(path))


def write_pdf(r: Resume, path: Path, font: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    pdfmetrics.registerFont(TTFont("Korean", str(font)))
    title = ParagraphStyle("t", fontName="Korean", fontSize=16, leading=22)
    sub = ParagraphStyle("s", fontName="Korean", fontSize=12, leading=18, spaceBefore=8)
    body = ParagraphStyle("b", fontName="Korean", fontSize=9, leading=13)
    story = [Paragraph(escape(f"이력서 - {r.title}"), title), Spacer(1, 6)]
    for heading, head, rows in sections(r):
        if not rows:
            continue
        story.append(Paragraph(escape(heading), sub))
        cells = [[Paragraph(escape(v), body) for v in row] for row in [head, *rows]]
        table = Table(cells, hAlign="LEFT")
        grid = ("GRID", (0, 0), (-1, -1), 0.5, colors.grey)
        table.setStyle(TableStyle([grid, ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke)]))
        story.append(table)

    SimpleDocTemplate(str(path), pagesize=A4, subject=NOTICE).build(story)


def write_hwpx(r: Resume, path: Path) -> None:
    """한글 표준 빈 문서 틀(python-hwpx)에 제목·표를 넣습니다. 손으로 만든 최소 구성은
    한글 프로그램에서 내용이 비어 보여서 라이브러리 틀을 씁니다."""
    from hwpx.document import HwpxDocument

    doc = HwpxDocument.new()
    doc.add_heading(f"이력서 - {r.title}", level=1)
    for heading, head, rows in sections(r):
        if not rows:
            continue
        doc.add_heading(heading, level=2)
        table = doc.add_table(1 + len(rows), len(head))
        for i, row in enumerate([head, *rows]):
            for j, value in enumerate(row):
                table.set_cell_text(i, j, value)
    doc.package.set_document_metadata(title=f"이력서 - {r.title}", subject=NOTICE)
    doc.save_to_path(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--font", type=Path, default=DEFAULT_FONT, help="PDF용 한글 TTF")
    parser.add_argument(
        "--today", type=date.fromisoformat, default=today_kst(), help="시연 입장일(YYYY-MM-DD)"
    )
    args = parser.parse_args()
    separations = {
        p["persona_id"]: at(args.today, p["procedure"]["separation_offset"])
        for p in read_seed_file()
    }
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for base in RESUMES:
        resume = with_separation(base, separations[base.persona])
        stem = OUT_DIR / f"{resume.persona}_resume"
        write_docx(resume, stem.with_suffix(".docx"))
        write_hwpx(resume, stem.with_suffix(".hwpx"))
        write_pdf(resume, stem.with_suffix(".pdf"), args.font)
        print(f"{resume.persona}: pdf·docx·hwpx")


if __name__ == "__main__":
    main()

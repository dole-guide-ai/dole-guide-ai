"""문서 첨부(pdf·docx·hwpx·hwp)에서 글자를 뽑습니다.

글자가 들어 있는 문서는 OCR 없이 서버 안에서 바로 읽으므로 원본이 외부 모델로 나가지 않습니다.
가리기(pii.mask)는 부르는 쪽이 합니다. 글자가 없는 스캔 pdf는 빈 글을 돌려주고, 부르는 쪽이
`pdf_page_images`로 쪽을 그림으로 바꿔 사진처럼 받아씁니다(preprocess).
"""

import io
import re
import struct
import zipfile
import zlib
from pathlib import Path
from xml.sax.saxutils import unescape

import olefile
import pypdfium2 as pdfium
from docx import Document
from pypdf import PdfReader

# 한 파일에서 대화로 넘기는 최대 글자 수. 경력증명서·이력서는 보통 1~2쪽이라 넉넉합니다.
MAX_DOCUMENT_CHARS = 4000
# docx·hwpx(zip)를 풀었을 때 크기 상한. 4MB 업로드가 수백 MB로 풀리는 압축 폭탄을 막습니다.
MAX_UNZIPPED_BYTES = 20 * 1024 * 1024
PDF_SIGNATURE = b"%PDF"
ZIP_SIGNATURE = b"PK\x03\x04"
HWPX_SECTION = re.compile(r"Contents/section\d+\.xml")
HWPX_TEXT = re.compile(r"<hp:t(?:\s[^>]*)?>([^<]*)</hp:t>")
OLE_SIGNATURE = b"\xd0\xcf\x11\xe0"
# 스캔 pdf에서 그림으로 바꿔 받아쓸 최대 쪽수와 배율(약 150dpi). 이력서·증명서는 보통 1~2쪽입니다.
MAX_SCANNED_PAGES = 3
PDF_RENDER_SCALE = 2
# 쪽 그림의 긴 변 상한(픽셀). 받아쓰기에서 어차피 1600px로 줄이므로 넉넉합니다.
MAX_RENDER_SIDE = 2400
# hwp 본문 영역(Section) 최대 개수. 이력서·증명서는 보통 1개입니다.
MAX_HWP_SECTIONS = 10

# 한글(HWP 5) 문서 구조: FileHeader의 속성 비트, 본문 레코드의 문단 글자 태그(HWPTAG_BEGIN+51)
HWP_HEADER_PROPERTIES = 36
HWP_COMPRESSED = 0x1
HWP_ENCRYPTED = 0x2
HWP_PARA_TEXT = 67
HWP_LONG_RECORD = 0xFFF
# 글자 사이의 제어 문자. 이 번호들은 8글자(16바이트) 자리를 차지합니다(표·그림 같은 개체, 탭).
HWP_WIDE_CONTROLS = frozenset(
    {1, 2, 3, 4, 5, 6, 7, 8, 9, 11, 12, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23}
)
HWP_WIDE_UNITS = 8
HWP_TAB = 9
HWP_LINE_BREAKS = frozenset({10, 13})


def _pdf_text(path: Path) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)


def _docx_text(path: Path) -> str:
    document = Document(str(path))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _hwpx_text(archive: zipfile.ZipFile) -> str:
    sections = sorted(n for n in archive.namelist() if HWPX_SECTION.fullmatch(n))
    texts = []
    for name in sections:
        xml = archive.read(name).decode("utf-8", errors="ignore")
        texts.append(" ".join(unescape(t) for t in HWPX_TEXT.findall(xml)))
    return "\n".join(texts)


def _inflate(data: bytes, budget: int) -> bytes:
    """압축된 본문을 풉니다. 남은 상한(budget)까지만 풀어 압축 폭탄을 막습니다."""
    inflater = zlib.decompressobj(-zlib.MAX_WBITS)
    return inflater.decompress(data, max(budget, 0))


def _hwp_records(data: bytes):
    """본문 레코드(태그, 내용)를 차례로."""
    offset = 0
    while offset + 4 <= len(data):
        (header,) = struct.unpack_from("<I", data, offset)
        offset += 4
        tag, size = header & 0x3FF, (header >> 20) & 0xFFF
        if size == HWP_LONG_RECORD:
            (size,) = struct.unpack_from("<I", data, offset)
            offset += 4
        yield tag, data[offset : offset + size]
        offset += size


def _hwp_paragraph(payload: bytes) -> str:
    """문단 글자(UTF-16LE). 개체·필드 같은 제어 문자는 건너뛰고 탭·줄바꿈만 남깁니다."""
    units = struct.unpack(f"<{len(payload) // 2}H", payload[: len(payload) // 2 * 2])
    out, i = [], 0
    while i < len(units):
        code = units[i]
        if code >= 32:
            out.append(chr(code))
        elif code == HWP_TAB:
            out.append("\t")
        elif code in HWP_LINE_BREAKS:
            out.append("\n")
        i += HWP_WIDE_UNITS if code in HWP_WIDE_CONTROLS else 1
    return "".join(out)


def _hwp_text(path: Path) -> str:
    """옛 한글(HWP 5, OLE 복합 문서) 본문. 암호가 걸린 문서는 읽지 않습니다."""
    with olefile.OleFileIO(str(path)) as ole:
        header = ole.openstream("FileHeader").read()
        (properties,) = struct.unpack_from("<I", header, HWP_HEADER_PROPERTIES)
        if properties & HWP_ENCRYPTED:
            return ""
        sections = sorted(
            (e for e in ole.listdir() if len(e) == 2 and e[0] == "BodyText"),
            key=lambda e: int(e[1].removeprefix("Section") or 0),
        )
        # 풀린 크기는 파일 전체 합으로 제한하고, 넘길 글자가 다 모이면 멈춥니다.
        texts, budget, chars = [], MAX_UNZIPPED_BYTES, 0
        for entry in sections[:MAX_HWP_SECTIONS]:
            data = ole.openstream(entry).read(budget)
            if properties & HWP_COMPRESSED:
                data = _inflate(data, budget)
            budget -= len(data)
            for tag, payload in _hwp_records(data):
                if tag == HWP_PARA_TEXT:
                    texts.append(_hwp_paragraph(payload))
                    chars += len(texts[-1])
            if budget <= 0 or chars >= MAX_DOCUMENT_CHARS:
                break
    return "\n".join(texts)


def pdf_page_images(path: str) -> list[bytes]:
    """pdf 앞 몇 쪽을 PNG로. 글자가 없는 스캔 pdf를 사진처럼 받아쓸 때 씁니다. 못 읽으면 빈 목록."""
    try:
        pdf = pdfium.PdfDocument(path)
    except Exception:  # 깨진 pdf
        return []
    pages = []
    try:
        for index in range(min(len(pdf), MAX_SCANNED_PAGES)):
            page = pdf[index]
            # 쪽 크기(포인트)가 아주 커도 긴 변이 MAX_RENDER_SIDE 픽셀을 넘지 않게 배율을 줄입니다.
            scale = min(PDF_RENDER_SCALE, MAX_RENDER_SIDE / max(*page.get_size(), 1))
            image = page.render(scale=scale).to_pil()
            buffer = io.BytesIO()
            image.save(buffer, "PNG")
            pages.append(buffer.getvalue())
    except Exception:  # 그리다 실패한 쪽부터는 버리고 그린 쪽까지만 씁니다.
        pass
    finally:
        pdf.close()
    return pages


def is_pdf(path: str) -> bool:
    with Path(path).open("rb") as f:
        return f.read(4).startswith(PDF_SIGNATURE)


def _zip_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        if sum(info.file_size for info in archive.infolist()) > MAX_UNZIPPED_BYTES:
            return ""
        names = archive.namelist()
        if "word/document.xml" in names:
            return _docx_text(path)
        if any(HWPX_SECTION.fullmatch(n) for n in names):
            return _hwpx_text(archive)
    return ""


def extract_text(path: str) -> str:
    """문서 글자. 형식은 확장자가 아니라 파일 첫 바이트로 정합니다. 못 읽으면 빈 글."""
    file = Path(path)
    head = file.read_bytes()[:4]
    try:
        if head.startswith(PDF_SIGNATURE):
            text = _pdf_text(file)
        elif head.startswith(ZIP_SIGNATURE):
            text = _zip_text(file)
        elif head.startswith(OLE_SIGNATURE):
            text = _hwp_text(file)
        else:
            return ""
    except Exception:  # 깨진 문서는 대화를 막지 않고 못 읽은 것으로 칩니다.
        return ""
    lines = (line.strip() for line in text.splitlines())
    return "\n".join(line for line in lines if line)[:MAX_DOCUMENT_CHARS]

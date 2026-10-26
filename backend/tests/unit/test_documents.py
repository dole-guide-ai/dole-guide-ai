"""첨부 문서 글자 뽑기(agent/input/documents.py). 파일은 테스트 안에서 만든 합성 값입니다."""

import zipfile

from docx import Document

from app.agent.input import documents


def minimal_pdf(text: str) -> bytes:
    """글자 한 줄짜리 최소 PDF(Helvetica). xref 위치를 계산해 넣어야 pypdf가 읽습니다."""
    stream = f"BT /F1 12 Tf 10 50 Td ({text}) Tj ET".encode()
    objects = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 100]/Contents 4 0 R"
        b"/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length %d>>stream\n" % len(stream) + stream + b"\nendstream",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj" % number + body + b"endobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer<</Size %d/Root 1 0 R>>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return bytes(out)


def test_pdf_text(tmp_path):
    path = tmp_path / "career.pdf"
    path.write_bytes(minimal_pdf("Hello OO Logistics"))
    assert "Hello OO Logistics" in documents.extract_text(str(path))


def test_docx_paragraphs_and_tables(tmp_path):
    path = tmp_path / "career.docx"
    doc = Document()
    doc.add_paragraph("경력증명서")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "회사명"
    table.cell(0, 1).text = "OO물류"
    doc.save(str(path))

    text = documents.extract_text(str(path))

    assert "경력증명서" in text
    assert "회사명 | OO물류" in text


def test_hwpx_section_text(tmp_path):
    path = tmp_path / "license.hwpx"
    xml = "<hs:sec><hp:p><hp:run><hp:t>자격증명: 지게차운전기능사</hp:t></hp:run></hp:p></hs:sec>"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("mimetype", "application/hwp+zip")
        archive.writestr("Contents/section0.xml", xml)

    assert documents.extract_text(str(path)) == "자격증명: 지게차운전기능사"


def test_unknown_or_broken_file_is_empty(tmp_path):
    hwp = tmp_path / "old.hwp"
    hwp.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 16)
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"%PDF-1.4 not really")

    assert documents.extract_text(str(hwp)) == ""
    assert documents.extract_text(str(broken)) == ""


def test_long_text_is_cut(tmp_path):
    path = tmp_path / "long.docx"
    doc = Document()
    doc.add_paragraph("가" * (documents.MAX_DOCUMENT_CHARS + 100))
    doc.save(str(path))

    assert len(documents.extract_text(str(path))) == documents.MAX_DOCUMENT_CHARS
